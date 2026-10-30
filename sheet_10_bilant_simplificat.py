#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Capitolul 10: Bilanțul Financiar Simplificat.
Comprimă patrimoniul în cele trei mase de activ (AIN = Active imobilizate nete, Disponibil,
AC fără disponibil) și cele două mase de pasiv (CP = Capitaluri permanente, DEX = Datorii pe termen scurt).
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from core_balanta import (
    BalantaData,
    add_heading,
    add_paragraph,
    fmt_money,
    fmt_ratio_pct,
    make_table,
    pick_variant,
    safe_div,
)
from sheet_03_bilant_financiar import BilantFinanciarResult
from sheet_07_echilibrul_financiar import EchilibruFinanciarResult


@dataclass
class BilantSimplificatResult:
    ain: Tuple[float, float, float]
    disponibil: Tuple[float, float, float]
    ac_fara_disp: Tuple[float, float, float]
    total_activ: Tuple[float, float, float]
    cp_perm: Tuple[float, float, float]
    dex: Tuple[float, float, float]
    total_pasiv: Tuple[float, float, float]


def compute_sheet_10_bilant_simplificat(
    bf: BilantFinanciarResult, ef: EchilibruFinanciarResult
) -> BilantSimplificatResult:
    return BilantSimplificatResult(
        ain=bf.active_imobilizate,
        disponibil=bf.disponibilitati,
        ac_fara_disp=ef.ac_fara_disp,
        total_activ=bf.total_activ,
        cp_perm=bf.capitaluri_permanente,
        dex=bf.datorii_termen_scurt,
        total_pasiv=bf.total_pasiv,
    )


def render_sheet_10_bilant_simplificat(
    doc: Document, bal: BalantaData, res: BilantSimplificatResult
) -> None:
    """Generează secțiunea Bilanțului Financiar Simplificat în documentul Word."""
    add_heading(doc, "10. Bilanțul Financiar Simplificat (Structura Sintetică Activ–Pasiv)", level=1)

    comp_headers = [
        "Masă Bilanțieră Simplificată (Activ / Pasiv)",
        "Perioada de raportare 0 (lei)",
        "Pond. P0",
        "Perioada de raportare 1 (lei)",
        "Pond. P1",
        "Perioada de raportare 2 (lei)",
        "Pond. P2",
    ]
    comp_align = [
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.RIGHT,
        WD_ALIGN_PARAGRAPH.RIGHT,
        WD_ALIGN_PARAGRAPH.RIGHT,
        WD_ALIGN_PARAGRAPH.RIGHT,
        WD_ALIGN_PARAGRAPH.RIGHT,
        WD_ALIGN_PARAGRAPH.RIGHT,
    ]
    comp_rows = [
        [
            "AIN — Active imobilizate nete (Activ)",
            fmt_money(res.ain[0]),
            fmt_ratio_pct(safe_div(res.ain[0], res.total_activ[0])),
            fmt_money(res.ain[1]),
            fmt_ratio_pct(safe_div(res.ain[1], res.total_activ[1])),
            fmt_money(res.ain[2]),
            fmt_ratio_pct(safe_div(res.ain[2], res.total_activ[2])),
        ],
        [
            "Disponibil — Trezorerie activă (Activ)",
            fmt_money(res.disponibil[0]),
            fmt_ratio_pct(safe_div(res.disponibil[0], res.total_activ[0])),
            fmt_money(res.disponibil[1]),
            fmt_ratio_pct(safe_div(res.disponibil[1], res.total_activ[1])),
            fmt_money(res.disponibil[2]),
            fmt_ratio_pct(safe_div(res.disponibil[2], res.total_activ[2])),
        ],
        [
            "AC fără disponibil — Stocuri + Creanțe (Activ)",
            fmt_money(res.ac_fara_disp[0]),
            fmt_ratio_pct(safe_div(res.ac_fara_disp[0], res.total_activ[0])),
            fmt_money(res.ac_fara_disp[1]),
            fmt_ratio_pct(safe_div(res.ac_fara_disp[1], res.total_activ[1])),
            fmt_money(res.ac_fara_disp[2]),
            fmt_ratio_pct(safe_div(res.ac_fara_disp[2], res.total_activ[2])),
        ],
        [
            "TOTAL ACTIV = TOTAL PASIV",
            fmt_money(res.total_activ[0]),
            "100,00%",
            fmt_money(res.total_activ[1]),
            "100,00%",
            fmt_money(res.total_activ[2]),
            "100,00%",
        ],
        [
            "CP — Capitaluri permanente (Pasiv)",
            fmt_money(res.cp_perm[0]),
            fmt_ratio_pct(safe_div(res.cp_perm[0], res.total_pasiv[0])),
            fmt_money(res.cp_perm[1]),
            fmt_ratio_pct(safe_div(res.cp_perm[1], res.total_pasiv[1])),
            fmt_money(res.cp_perm[2]),
            fmt_ratio_pct(safe_div(res.cp_perm[2], res.total_pasiv[2])),
        ],
        [
            "DEX — Datorii pe termen scurt / exploatare (Pasiv)",
            fmt_money(res.dex[0]),
            fmt_ratio_pct(safe_div(res.dex[0], res.total_pasiv[0])),
            fmt_money(res.dex[1]),
            fmt_ratio_pct(safe_div(res.dex[1], res.total_pasiv[1])),
            fmt_money(res.dex[2]),
            fmt_ratio_pct(safe_div(res.dex[2], res.total_pasiv[2])),
        ],
    ]
    make_table(
        doc,
        headers=comp_headers,
        rows=comp_rows,
        col_widths=[5.6, 2.3, 1.7, 2.3, 1.7, 2.3, 1.7],
        align_map=comp_align,
        total_rows=[3],
        font_size=7.8,
    )

    p_ain2 = safe_div(res.ain[2], res.total_activ[2])
    p_disp2 = safe_div(res.disponibil[2], res.total_activ[2])
    p_acf2 = safe_div(res.ac_fara_disp[2], res.total_activ[2])
    p_cp2 = safe_div(res.cp_perm[2], res.total_pasiv[2])
    p_dex2 = safe_div(res.dex[2], res.total_pasiv[2])

    s10_msg = pick_variant(
        bal,
        "s10_struct",
        [
            f"Sinteză structurală simplificată: La nivelul perioadei de raportare 2, din totalul activului de {fmt_money(res.total_activ[2])} lei, imobilizările nete (AIN) însumează {fmt_money(res.ain[2])} lei ({fmt_ratio_pct(p_ain2)}), stocurile și creanțele (AC fără disponibil) reprezintă {fmt_money(res.ac_fara_disp[2])} lei ({fmt_ratio_pct(p_acf2)}), iar trezoreria activă (Disponibilul) deține {fmt_money(res.disponibil[2])} lei ({fmt_ratio_pct(p_disp2)}). Pe latura pasivului, finanțarea este asigurată în proporție de {fmt_ratio_pct(p_cp2)} din capitaluri permanente ({fmt_money(res.cp_perm[2])} lei) și {fmt_ratio_pct(p_dex2)} din datorii pe termen scurt ({fmt_money(res.dex[2])} lei).",
            f"Arhitectura bilanțului simplificat în perioada de raportare 2 evidențiază echilibrul dintre utilizările patrimoniale — compuse din active imobilizate nete de {fmt_money(res.ain[2])} lei ({fmt_ratio_pct(p_ain2)}), active circulante non-monetare de {fmt_money(res.ac_fara_disp[2])} lei ({fmt_ratio_pct(p_acf2)}) și disponibilități bănești de {fmt_money(res.disponibil[2])} lei ({fmt_ratio_pct(p_disp2)}) — și resursele de finanțare, constituite din capitaluri permanente ({fmt_money(res.cp_perm[2])} lei, {fmt_ratio_pct(p_cp2)}) și datorii de exploatare pe termen scurt ({fmt_money(res.dex[2])} lei, {fmt_ratio_pct(p_dex2)}).",
            f"Gruparea sintetică a posturilor bilanțiere în perioada de raportare 2 arată că resursele permanente ({fmt_money(res.cp_perm[2])} lei, reprezentând {fmt_ratio_pct(p_cp2)} din pasiv) împreună cu datoriile curente ({fmt_money(res.dex[2])} lei, {fmt_ratio_pct(p_dex2)}) susțin un activ format din imobilizări nete ({fmt_ratio_pct(p_ain2)}), stocuri și creanțe ({fmt_ratio_pct(p_acf2)}) și lichidități imediate ({fmt_ratio_pct(p_disp2)}).",
        ],
    )
    add_paragraph(doc, s10_msg)
