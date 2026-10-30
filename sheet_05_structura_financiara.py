#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Capitolul 5: Analiza Structurii Financiare a Pasivului.
Analizează arhitectura surselor de finanțare prin optica apartenenței fondurilor
(Total Pasiv 1 = Capitaluri proprii + Datorii totale) și optica stabilității în timp
(Total Pasiv 2 = Capitaluri permanente + Datorii pe termen scurt).
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Tuple

from docx import Document

from core_balanta import (
    BalantaData,
    add_heading,
    add_paragraph,
    calc_delta,
    calc_index,
    describe_trajectory,
    fmt_money,
    fmt_pct,
    fmt_ratio_pct,
    make_table,
    pick_variant,
    safe_div,
)
from sheet_03_bilant_financiar import BilantFinanciarResult


@dataclass
class StructuraFinanciaraResult:
    cp: Tuple[float, float, float]
    dtl: Tuple[float, float, float]
    dts: Tuple[float, float, float]
    cperm: Tuple[float, float, float]
    dtot: Tuple[float, float, float]
    pasiv1: Tuple[float, float, float]
    pasiv2: Tuple[float, float, float]


def compute_sheet_05_structura_financiara(bf: BilantFinanciarResult) -> StructuraFinanciaraResult:
    cp = bf.capitaluri_proprii
    dtl = bf.datorii_termen_lung
    dts = bf.datorii_termen_scurt
    cperm = bf.capitaluri_permanente
    dtot = tuple(round(dtl[p] + dts[p], 2) for p in range(3))
    pasiv1 = tuple(round(cp[p] + dtot[p], 2) for p in range(3))
    pasiv2 = tuple(round(cperm[p] + dts[p], 2) for p in range(3))
    return StructuraFinanciaraResult(
        cp=cp,
        dtl=dtl,
        dts=dts,
        cperm=cperm,
        dtot=dtot,
        pasiv1=pasiv1,
        pasiv2=pasiv2,
    )


def render_sheet_05_structura_financiara(
    doc: Document, bal: BalantaData, sf: StructuraFinanciaraResult
) -> None:
    """Generează secțiunea Structurii Financiare a Pasivului în documentul Word."""
    add_heading(doc, "5. Analiza Structurii Financiare a Pasivului", level=1)

    items = [
        ("1", "Capitaluri proprii (CP)", sf.cp, False),
        ("2", "Datorii pe termen lung (DTL)", sf.dtl, False),
        ("3", "Datorii pe termen scurt (DTS)", sf.dts, False),
        ("4", "Capitaluri permanente (CPERM = CP + DTL)", sf.cperm, True),
        ("5", "Datorii totale (DT = DTL + DTS)", sf.dtot, True),
        ("6", "Total Pasiv 1 (Capitaluri proprii + Datorii totale)", sf.pasiv1, True),
        ("7", "Total Pasiv 2 (Capitaluri permanente + Datorii pe termen scurt)", sf.pasiv2, True),
    ]

    headers = [
        "Nr.",
        "Agregat de Structură Financiară (lei)",
        "Perioada de raportare 0",
        "Perioada de raportare 1",
        "Perioada de raportare 2",
        "Δ (P1 - P0)",
        "Δ (P2 - P1)",
        "I (P1 / P0)",
        "I (P2 / P1)",
    ]
    rows = []
    tot_idx = []
    for i, (nr, name, vals, is_tot) in enumerate(items):
        if is_tot:
            tot_idx.append(i)
        d = calc_delta(vals)
        idx = calc_index(vals)
        rows.append([
            nr,
            name,
            fmt_money(vals[0]),
            fmt_money(vals[1]),
            fmt_money(vals[2]),
            fmt_money(d[0]),
            fmt_money(d[1]),
            fmt_ratio_pct(idx[0]),
            fmt_ratio_pct(idx[1]),
        ])

    make_table(
        doc,
        headers=headers,
        rows=rows,
        col_widths=[0.7, 5.1, 1.9, 1.9, 1.9, 1.7, 1.7, 1.45, 1.45],
        total_rows=tot_idx,
        font_size=7.7,
    )

    traj_cp = describe_trajectory(sf.cp, bal, "s05_cp", "capitalurile proprii")
    traj_dtot = describe_trajectory(sf.dtot, bal, "s05_dtot", "datoriile totale (DTL + DTS)")
    pond_cp2 = safe_div(sf.cp[2], sf.pasiv1[2])

    if pond_cp2 >= 0.50:
        aut_msg = pick_variant(
            bal,
            "s05_aut_high",
            [
                f"În perioada de raportare 2, autofinanțarea deține rolul determinant ({fmt_ratio_pct(pond_cp2)} din pasivul total), conferind societății o independență financiară solidă față de creditori.",
                f"Preponderența resurselor proprii în pasivul perioadei de raportare 2 ({fmt_ratio_pct(pond_cp2)}) atestă o structură de finanțare conservatoare și rezilientă la șocuri pe piața creditului.",
                f"Capitalurile proprii acoperă {fmt_ratio_pct(pond_cp2)} din totalul resurselor în perioada de raportare 2, devansând datoriile totale și consolidând autonomia decizională a firmei.",
            ],
        )
    else:
        aut_msg = pick_variant(
            bal,
            "s05_aut_low",
            [
                f"În perioada de raportare 2, resursele atrase și împrumutate dețin {fmt_ratio_pct(1.0 - pond_cp2)} din pasivul total (capitalurile proprii reprezentând {fmt_ratio_pct(pond_cp2)}), finanțarea activității bazându-se prioritar pe creditul comercial și bancar.",
                f"Ponderea capitalurilor proprii în pasivul perioadei de raportare 2 este de {fmt_ratio_pct(pond_cp2)}, ceea ce reflectă o utilizare activă a resurselor temporare și la termen ({fmt_money(sf.dtot[2])} lei) în susținerea ciclului de exploatare.",
            ],
        )

    add_paragraph(
        doc,
        f"Dinamica structurii financiare: Pe parcursul celor trei perioade de raportare, {traj_cp}, în timp ce {traj_dtot}. {aut_msg}",
    )
