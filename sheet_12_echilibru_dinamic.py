#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Capitolul 12: Diagnosticul Echilibrului Financiar Dinamic (ΔFR, ΔNFR, ΔTN).
Analizează variațiile indicatorilor de echilibru între Perioada de raportare 0 și Perioada de raportare 1,
respectiv între Perioada de raportare 1 și Perioada de raportare 2.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Tuple

from docx import Document

from core_balanta import (
    BalantaData,
    add_heading,
    add_paragraph,
    calc_delta,
    fmt_money,
    make_table,
    pick_variant,
)
from sheet_11_echilibru_static import EchilibruStaticResult


@dataclass
class EchilibruDinamicResult:
    d_fr: Tuple[float, float]
    d_nfr: Tuple[float, float]
    d_tn: Tuple[float, float]
    interp_d_fr: Tuple[str, str]
    interp_d_nfr: Tuple[str, str]
    interp_d_tn: Tuple[str, str]


def compute_sheet_12_echilibru_dinamic(es: EchilibruStaticResult) -> EchilibruDinamicResult:
    d_fr = calc_delta(es.fr)
    d_nfr = calc_delta(es.nfr)
    d_tn = calc_delta(es.tn)

    interp_fr = tuple("EFTL" if d_fr[i] >= 0 else "DFTL" for i in (0, 1))
    interp_nfr = tuple("EFTS" if d_fr[i] >= d_nfr[i] else "NFCE" for i in (0, 1))
    interp_tn = tuple("EFC" if d_tn[i] >= 0 else "DFT" for i in (0, 1))

    return EchilibruDinamicResult(
        d_fr=d_fr,
        d_nfr=d_nfr,
        d_tn=d_tn,
        interp_d_fr=interp_fr,
        interp_d_nfr=interp_nfr,
        interp_d_tn=interp_tn,
    )


def render_sheet_12_echilibru_dinamic(
    doc: Document, bal: BalantaData, res: EchilibruDinamicResult
) -> None:
    """Generează secțiunea Echilibrului Financiar Dinamic în documentul Word."""
    add_heading(doc, "12. Diagnosticul Echilibrului Financiar Dinamic (ΔFR, ΔNFR, ΔTN)", level=1)

    headers = [
        "Indicator Dinamic (Fluxul Perioadei)",
        "Modificare P1 - P0 (lei)",
        "Diagnostic P1 - P0",
        "Modificare P2 - P1 (lei)",
        "Diagnostic P2 - P1",
        "Modificare Totală P2 - P0 (lei)",
    ]
    rows = [
        [
            "ΔFR (Modificarea Fondului de rulment)",
            fmt_money(res.d_fr[0]),
            res.interp_d_fr[0],
            fmt_money(res.d_fr[1]),
            res.interp_d_fr[1],
            fmt_money(res.d_fr[0] + res.d_fr[1]),
        ],
        [
            "ΔNFR (Modificarea Necesarului de fond de rulment)",
            fmt_money(res.d_nfr[0]),
            res.interp_d_nfr[0],
            fmt_money(res.d_nfr[1]),
            res.interp_d_nfr[1],
            fmt_money(res.d_nfr[0] + res.d_nfr[1]),
        ],
        [
            "ΔTN = ΔFR - ΔNFR (Modificarea Trezoreriei nete)",
            fmt_money(res.d_tn[0]),
            res.interp_d_tn[0],
            fmt_money(res.d_tn[1]),
            res.interp_d_tn[1],
            fmt_money(res.d_tn[0] + res.d_tn[1]),
        ],
    ]
    make_table(
        doc,
        headers=headers,
        rows=rows,
        col_widths=[5.2, 2.6, 2.2, 2.6, 2.2, 2.8],
        total_rows=[2],
        font_size=7.8,
    )

    tot_dfr = res.d_fr[0] + res.d_fr[1]
    tot_dnfr = res.d_nfr[0] + res.d_nfr[1]
    tot_dtn = res.d_tn[0] + res.d_tn[1]

    s12_msg = pick_variant(
        bal,
        "s12_dyn",
        [
            f"Sinteză dinamică a fluxurilor: Între perioada de raportare 0 și perioada de raportare 1, variațiile au fost ΔFR = {fmt_money(res.d_fr[0])} lei ({res.interp_d_fr[0]}), ΔNFR = {fmt_money(res.d_nfr[0])} lei ({res.interp_d_nfr[0]}) și ΔTN = {fmt_money(res.d_tn[0])} lei ({res.interp_d_tn[0]}). În intervalul dintre perioada de raportare 1 și perioada de raportare 2, fluxurile consemnează ΔFR = {fmt_money(res.d_fr[1])} lei ({res.interp_d_fr[1]}), ΔNFR = {fmt_money(res.d_nfr[1])} lei ({res.interp_d_nfr[1]}) și ΔTN = {fmt_money(res.d_tn[1])} lei ({res.interp_d_tn[1]}), cumulând pe întregul orizont P2 - P0 o modificare de {fmt_money(tot_dfr)} lei a Fondului de rulment și de {fmt_money(tot_dtn)} lei a Trezoreriei nete.",
            f"Analiza variațiilor de echilibru arată că în primul subinterval (P1 - P0) fluxul net de trezorerie a fost ΔTN = {fmt_money(res.d_tn[0])} lei (pe fondul ΔFR = {fmt_money(res.d_fr[0])} lei și ΔNFR = {fmt_money(res.d_nfr[0])} lei), în timp ce în al doilea subinterval (P2 - P1) s-a înregistrat ΔFR = {fmt_money(res.d_fr[1])} lei ({res.interp_d_fr[1]}), ΔNFR = {fmt_money(res.d_nfr[1])} lei ({res.interp_d_nfr[1]}) și ΔTN = {fmt_money(res.d_tn[1])} lei ({res.interp_d_tn[1]}). Cumulat între perioada 0 și perioada 2, ΔFR este de {fmt_money(tot_dfr)} lei, ΔNFR de {fmt_money(tot_dnfr)} lei, iar ΔTN de {fmt_money(tot_dtn)} lei.",
        ],
    )
    add_paragraph(doc, s12_msg)
