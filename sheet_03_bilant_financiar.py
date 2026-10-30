#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Capitolul 3: Bilanțul Financiar Patrimonial (Retratat pe Criterii de Lichiditate–Exigibilitate).
Aplică retratările patrimoniale asupra Bilanțului Contabil:
  - Creanțe financiare = Creanțe bilanțiere (rd. 06) + Cheltuieli în avans (rd. 10, ct. 471)
  - Capitaluri proprii financiare = Capitaluri proprii bilanțiere (rd. 39) + Provizioane (rd. 15, ct. 151)
  - Datorii pe termen scurt financiare = Datorii sub un an (rd. 11) + Venituri în avans (rd. 16, ct. 475 + 472)
  - Capitaluri permanente = Capitaluri proprii financiare + Datorii pe termen lung (rd. 14)
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Dict, Tuple

from docx import Document

from core_balanta import (
    BalantaData,
    add_heading,
    add_paragraph,
    calc_delta,
    calc_index,
    describe_trajectory,
    fmt_money,
    fmt_ratio_pct,
    make_table,
    pick_variant,
    safe_div,
)
from sheet_01_bilant import BilantResult


@dataclass
class BilantFinanciarResult:
    rows: Dict[int, Tuple[float, float, float]] = field(default_factory=dict)

    def r(self, idx: int) -> Tuple[float, float, float]:
        return self.rows.get(idx, (0.0, 0.0, 0.0))

    @property
    def imobilizari_necorporale(self) -> Tuple[float, float, float]:
        return self.r(1)

    @property
    def imobilizari_corporale(self) -> Tuple[float, float, float]:
        return self.r(2)

    @property
    def imobilizari_financiare(self) -> Tuple[float, float, float]:
        return self.r(3)

    @property
    def active_imobilizate(self) -> Tuple[float, float, float]:
        return self.r(4)

    @property
    def stocuri(self) -> Tuple[float, float, float]:
        return self.r(5)

    @property
    def creante(self) -> Tuple[float, float, float]:
        return self.r(6)

    @property
    def investitii_ts(self) -> Tuple[float, float, float]:
        return self.r(7)

    @property
    def disponibilitati(self) -> Tuple[float, float, float]:
        return self.r(8)

    @property
    def active_circulante(self) -> Tuple[float, float, float]:
        return self.r(9)

    @property
    def total_activ(self) -> Tuple[float, float, float]:
        return self.r(10)

    @property
    def capitaluri_proprii(self) -> Tuple[float, float, float]:
        return self.r(11)

    @property
    def datorii_termen_lung(self) -> Tuple[float, float, float]:
        return self.r(12)

    @property
    def capitaluri_permanente(self) -> Tuple[float, float, float]:
        return self.r(13)

    @property
    def datorii_termen_scurt(self) -> Tuple[float, float, float]:
        return self.r(14)

    @property
    def total_pasiv(self) -> Tuple[float, float, float]:
        return self.r(15)


def compute_sheet_03_bilant_financiar(bil: BilantResult) -> BilantFinanciarResult:
    """Calculează posturile Bilanțului Financiar Patrimonial pentru perioadele de raportare 0, 1 și 2."""
    res = BilantFinanciarResult()

    r1 = bil.rd(1)
    r2 = bil.rd(2)
    r3 = bil.rd(3)
    r4 = bil.rd(4)
    r5 = bil.rd(5)
    c_bil = bil.rd(6)
    ch_av = bil.rd(10)
    r6 = tuple(round(c_bil[p] + ch_av[p], 2) for p in range(3))
    r7 = bil.rd(7)
    r8 = bil.rd(8)
    r9 = tuple(round(r5[p] + r6[p] + r7[p] + r8[p], 2) for p in range(3))
    r10 = tuple(round(r4[p] + r9[p], 2) for p in range(3))

    cp_bil = bil.rd(39)
    prov = bil.rd(15)
    r11 = tuple(round(cp_bil[p] + prov[p], 2) for p in range(3))
    r12 = bil.rd(14)
    r13 = tuple(round(r11[p] + r12[p], 2) for p in range(3))
    dts_bil = bil.rd(11)
    ven_av = bil.rd(16)
    r14 = tuple(round(dts_bil[p] + ven_av[p], 2) for p in range(3))
    r15 = tuple(round(r13[p] + r14[p], 2) for p in range(3))

    res.rows = {
        1: r1, 2: r2, 3: r3, 4: r4, 5: r5,
        6: r6, 7: r7, 8: r8, 9: r9, 10: r10,
        11: r11, 12: r12, 13: r13, 14: r14, 15: r15,
    }
    return res


def render_sheet_03_bilant_financiar(
    doc: Document, bal: BalantaData, bil: BilantResult, bf: BilantFinanciarResult
) -> None:
    """Generează secțiunea Bilanțului Financiar Patrimonial în documentul Word."""
    add_heading(doc, "3. Bilanțul Financiar Patrimonial (Retratat pe Criterii de Lichiditate–Exigibilitate)", level=1)

    labels = [
        (1, "Imobilizări necorporale", False),
        (2, "Imobilizări corporale", False),
        (3, "Imobilizări financiare", False),
        (4, "ACTIVE IMOBILIZATE (rd. 1 + 2 + 3)", True),
        (5, "Stocuri", False),
        (6, "Creanțe retratate (Creanțe bilanț + Cheltuieli în avans ct. 471)", False),
        (7, "Investiții financiare pe termen scurt", False),
        (8, "Disponibilități bănești", False),
        (9, "ACTIVE CIRCULANTE (rd. 5 + 6 + 7 + 8)", True),
        (10, "TOTAL ACTIV FINANCIAR (rd. 4 + 9)", True),
        (11, "Capitaluri proprii (CP bilanț + Provizioane ct. 151)", False),
        (12, "Datorii pe termen lung (peste 1 an)", False),
        (13, "CAPITALURI PERMANENTE (rd. 11 + 12)", True),
        (14, "Datorii pe termen scurt (Datorii < 1 an + Venituri în avans ct. 472 + 475)", False),
        (15, "TOTAL PASIV FINANCIAR (rd. 13 + 14)", True),
    ]

    headers = [
        "Nr.",
        "Elemente Bilanț Financiar (lei)",
        "Perioada de raportare 0",
        "Perioada de raportare 1",
        "Perioada de raportare 2",
        "Δ (P1 - P0)",
        "Δ (P2 - P1)",
        "I (P2 / P1)",
    ]
    rows_data = []
    total_indices = []
    for i, (code, name, is_tot) in enumerate(labels):
        if is_tot:
            total_indices.append(i)
        vals = bf.r(code)
        d = calc_delta(vals)
        idx = calc_index(vals)
        rows_data.append([
            str(code),
            name,
            fmt_money(vals[0]),
            fmt_money(vals[1]),
            fmt_money(vals[2]),
            fmt_money(d[0]),
            fmt_money(d[1]),
            fmt_ratio_pct(idx[1]),
        ])

    make_table(
        doc,
        headers=headers,
        rows=rows_data,
        col_widths=[0.8, 5.8, 2.0, 2.0, 2.0, 1.8, 1.8, 1.6],
        total_rows=total_indices,
        font_size=7.8,
    )

    ch_av2 = bil.rd(10)[2]
    ven_av2 = bil.rd(16)[2]
    cperm_cov = safe_div(bf.capitaluri_permanente[2], bf.active_imobilizate[2])
    traj_cperm = describe_trajectory(bf.capitaluri_permanente, bal, "s03_cperm", "capitalurile permanente")

    retrat_det = []
    if ch_av2 > 0:
        retrat_det.append(f"includerea cheltuielilor în avans ({fmt_money(ch_av2)} lei) în creanțele retratate")
    if ven_av2 > 0:
        retrat_det.append(f"asimilarea veniturilor în avans ({fmt_money(ven_av2)} lei) în datoriile pe termen scurt")
    retrat_txt = (
        f"Prin {' și '.join(retrat_det)}, b" if retrat_det else "După reclasificarea posturilor pe criterii de lichiditate și exigibilitate, b"
    )

    if cperm_cov >= 1.0:
        cov_msg = pick_variant(
            bal,
            "s03_cov_ok",
            [
                f"Capitalurile permanente ({fmt_money(bf.capitaluri_permanente[2])} lei în perioada de raportare 2) acoperă de {fmt_num_cov(cperm_cov)} ori activele imobilizate ({fmt_money(bf.active_imobilizate[2])} lei), degajând o marjă de siguranță care finanțează o parte substanțială din activele circulante ({fmt_money(bf.active_circulante[2])} lei).",
                f"Respectarea regulii de aur a finanțării este confirmată în perioada de raportare 2: resursele permanente ({fmt_money(bf.capitaluri_permanente[2])} lei, {fmt_ratio_pct(safe_div(bf.capitaluri_permanente[2], bf.total_pasiv[2]))} din pasiv) depășesc nevoile permanente imobilizate ({fmt_money(bf.active_imobilizate[2])} lei), în timp ce datoriile pe termen scurt însumează {fmt_money(bf.datorii_termen_scurt[2])} lei.",
                f"În perioada de raportare 2, {traj_cperm}, asigurând finanțarea integrală a activelor imobilizate ({fmt_money(bf.active_imobilizate[2])} lei) și lăsând un excedent stabil pentru acoperirea activelor circulante ({fmt_money(bf.active_circulante[2])} lei).",
            ],
        )
    else:
        cov_msg = (
            f"În perioada de raportare 2, capitalurile permanente ({fmt_money(bf.capitaluri_permanente[2])} lei) acoperă doar {fmt_ratio_pct(cperm_cov)} din activele imobilizate ({fmt_money(bf.active_imobilizate[2])} lei), diferența fiind finanțată din datorii pe termen scurt ({fmt_money(bf.datorii_termen_scurt[2])} lei)."
        )

    add_paragraph(
        doc,
        f"Diagnostic după retratarea financiară: {retrat_txt}ilanțul financiar din perioada de raportare 2 aliniază un activ imobilizat de "
        f"{fmt_money(bf.active_imobilizate[2])} lei ({fmt_ratio_pct(safe_div(bf.active_imobilizate[2], bf.total_activ[2]))}) și un activ "
        f"circulant de {fmt_money(bf.active_circulante[2])} lei ({fmt_ratio_pct(safe_div(bf.active_circulante[2], bf.total_activ[2]))}). {cov_msg}",
    )


def fmt_num_cov(val: float) -> str:
    return f"{val:.2f}".replace(".", ",")
