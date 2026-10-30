#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Capitolul 13: Analiza Soldurilor Intermediare de Gestiune (Cascada SIG).
Construiește cascada completă a Soldurilor Intermediare de Gestiune (Marja comercială,
Producția exercițiului, Valoarea adăugată, Excedentul brut de exploatare, Rezultatul exploatării,
Rezultatul curent, Rezultatul net) pe cele trei perioade de raportare (0, 1 și 2).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from docx import Document

from core_balanta import (
    BalantaData,
    add_heading,
    add_paragraph,
    calc_delta,
    calc_growth_rate,
    calc_index,
    calc_ritm_mediu,
    describe_trajectory,
    fmt_money,
    fmt_ratio_pct,
    make_table,
    pick_variant,
    safe_div,
)
from sheet_02_cpp import CPPResult


@dataclass
class SIGRowData:
    name: str
    vals: Tuple[float, float, float]
    delta: Tuple[float, float]
    index: Tuple[Optional[float], Optional[float]]
    growth: Tuple[Optional[float], Optional[float]]
    ritm_mediu: Optional[float]
    is_marja: bool = False


@dataclass
class SIGResult:
    rows: List[SIGRowData] = field(default_factory=list)
    by_key: Dict[str, Tuple[float, float, float]] = field(default_factory=dict)

    @property
    def marja_comerciala(self) -> Tuple[float, float, float]:
        return self.by_key.get("mc", (0.0, 0.0, 0.0))

    @property
    def productia_exercitiului(self) -> Tuple[float, float, float]:
        return self.by_key.get("qe", (0.0, 0.0, 0.0))

    @property
    def consumuri_terti(self) -> Tuple[float, float, float]:
        return self.by_key.get("ct", (0.0, 0.0, 0.0))

    @property
    def valoarea_adaugata(self) -> Tuple[float, float, float]:
        return self.by_key.get("va", (0.0, 0.0, 0.0))

    @property
    def ebe(self) -> Tuple[float, float, float]:
        return self.by_key.get("ebe", (0.0, 0.0, 0.0))

    @property
    def rezultat_exploatare(self) -> Tuple[float, float, float]:
        return self.by_key.get("re", (0.0, 0.0, 0.0))

    @property
    def rezultat_curent(self) -> Tuple[float, float, float]:
        return self.by_key.get("rc", (0.0, 0.0, 0.0))

    @property
    def rezultat_net(self) -> Tuple[float, float, float]:
        return self.by_key.get("rn", (0.0, 0.0, 0.0))


def compute_sheet_13_sig(cpp: CPPResult) -> SIGResult:
    """Calculează cascada Soldurilor Intermediare de Gestiune (SIG)."""
    res = SIGResult()

    vanz_marf = cpp.rd(3)
    cost_marf = tuple(round(cpp.rd(16)[p] - cpp.rd(17)[p], 2) for p in range(3))
    mc = tuple(round(vanz_marf[p] - cost_marf[p], 2) for p in range(3))

    prod_vand = tuple(round(cpp.rd(2)[p] - cpp.rd(4)[p], 2) for p in range(3))
    prod_stoc = tuple(round(cpp.rd(7)[p] - cpp.rd(8)[p], 2) for p in range(3))
    prod_imob = cpp.rd(9)
    qe = tuple(round(prod_vand[p] + prod_stoc[p] + prod_imob[p], 2) for p in range(3))

    cons_terti = tuple(
        round(cpp.rd(13)[p] + cpp.rd(14)[p] + cpp.rd(15)[p] + cpp.rd(28)[p], 2) for p in range(3)
    )

    va = tuple(round(mc[p] + qe[p] - cons_terti[p], 2) for p in range(3))

    subv_expl = cpp.rd(6)
    imp_taxe = cpp.rd(29)
    ch_pers = cpp.rd(18)
    ebe = tuple(round(va[p] + subv_expl[p] - imp_taxe[p] - ch_pers[p], 2) for p in range(3))

    reluari_expl = tuple(round(cpp.rd(23)[p] + cpp.rd(26)[p] + cpp.rd(34)[p], 2) for p in range(3))
    alte_ven_expl = cpp.rd(10)
    amort_calc = tuple(round(cpp.rd(22)[p] + cpp.rd(25)[p] + cpp.rd(33)[p], 2) for p in range(3))
    alte_ch_expl = tuple(round(cpp.rd(30)[p] + cpp.rd(31)[p], 2) for p in range(3))
    re = tuple(
        round(ebe[p] + reluari_expl[p] + alte_ven_expl[p] - amort_calc[p] - alte_ch_expl[p], 2)
        for p in range(3)
    )

    ven_fin = cpp.rd(45)
    ch_fin = cpp.rd(52)
    rc = tuple(round(re[p] + ven_fin[p] - ch_fin[p], 2) for p in range(3))

    rez_extr = tuple(round(cpp.rd(59)[p] - cpp.rd(60)[p], 2) for p in range(3))
    imp_prof = tuple(round(cpp.rd(65)[p] + cpp.rd(66)[p], 2) for p in range(3))
    rn = tuple(round(rc[p] + rez_extr[p] - imp_prof[p], 2) for p in range(3))

    res.by_key = {
        "mc": mc,
        "qe": qe,
        "ct": cons_terti,
        "va": va,
        "ebe": ebe,
        "re": re,
        "rc": rc,
        "rn": rn,
    }

    raw_rows = [
        ("Vânzări de mărfuri", vanz_marf, False),
        ("- Costul mărfurilor vândute", cost_marf, False),
        ("= Marja comercială", mc, True),
        ("Producția vândută (netă de reduceri comerciale)", prod_vand, False),
        ("± Producția stocată", prod_stoc, False),
        ("± Producția imobilizată", prod_imob, False),
        ("= Producția exercițiului", qe, True),
        ("Marja comercială", mc, False),
        ("+ Producția exercițiului", qe, False),
        ("- Consumuri provenite de la terți", cons_terti, False),
        ("= Valoarea adăugată (1)", va, True),
        ("+ Subvenții de exploatare aferente CA", subv_expl, False),
        ("- Impozite, taxe și vărsăminte asimilate", imp_taxe, False),
        ("- Cheltuieli cu personalul", ch_pers, False),
        ("= Excedent brut de exploatare / Insuficiență brută (2)", ebe, True),
        ("+ Reluări asupra ajustărilor pentru depreciere privind exploatarea", reluari_expl, False),
        ("+ Alte venituri din exploatare (subvenții investiții și proiecte ct. 758)", alte_ven_expl, False),
        ("- Amortizări și ajustări pentru depreciere de exploatare calculate", amort_calc, False),
        ("- Alte cheltuieli de exploatare", alte_ch_expl, False),
        ("= Rezultatul exploatării (3)", re, True),
        ("+ Venituri financiare", ven_fin, False),
        ("- Cheltuieli financiare", ch_fin, False),
        ("= Rezultatul curent / brut (4)", rc, True),
        ("+ Rezultatul extraordinar", rez_extr, False),
        ("- Impozitul pe profit", imp_prof, False),
        ("= Rezultatul net al exercițiului (5)", rn, True),
    ]

    for name, vals, is_m in raw_rows:
        res.rows.append(
            SIGRowData(
                name=name,
                vals=vals,
                delta=calc_delta(vals),
                index=calc_index(vals),
                growth=calc_growth_rate(vals),
                ritm_mediu=calc_ritm_mediu(vals),
                is_marja=is_m,
            )
        )

    return res


def render_sheet_13_sig(doc: Document, bal: BalantaData, res: SIGResult) -> None:
    """Generează secțiunea Soldurilor Intermediare de Gestiune (SIG) într-un tabel sintetic unificat."""
    add_heading(doc, "13. Analiza Soldurilor Intermediare de Gestiune (Cascada SIG)", level=1)

    headers = [
        "Element / Sold Intermediar de Gestiune",
        "Perioada de raportare 0",
        "Perioada de raportare 1",
        "Perioada de raportare 2",
        "Δ (P1 - P0)",
        "Δ (P2 - P1)",
        "Ritm P2/P1",
        "Ritm mediu",
    ]
    rows_table = []
    tot_idx = []
    # Includem toate marjele SIG si elementele constitutive cu valori nenule
    for r in res.rows:
        if not r.is_marja and all(abs(v) < 1e-6 for v in r.vals):
            continue
        if r.is_marja:
            tot_idx.append(len(rows_table))
        rows_table.append([
            r.name,
            fmt_money(r.vals[0]),
            fmt_money(r.vals[1]),
            fmt_money(r.vals[2]),
            fmt_money(r.delta[0]),
            fmt_money(r.delta[1]),
            fmt_ratio_pct(r.growth[1], plus_sign=True),
            fmt_ratio_pct(r.ritm_mediu),
        ])

    make_table(
        doc,
        headers=headers,
        rows=rows_table,
        col_widths=[5.2, 1.9, 1.9, 1.9, 1.8, 1.8, 1.65, 1.65],
        total_rows=tot_idx,
        font_size=7.6,
    )

    traj_va = describe_trajectory(res.valoarea_adaugata, bal, "s13_va", "Valoarea adăugată (VA)")
    traj_ebe = describe_trajectory(res.ebe, bal, "s13_ebe", "Excedentul brut al exploatării (EBE)")

    mc2 = res.marja_comerciala[2]
    pe2 = res.productia_exercitiului[2]
    va2 = res.valoarea_adaugata[2]
    ebe2 = res.ebe[2]
    re2 = res.rezultat_exploatare[2]
    rc2 = res.rezultat_curent[2]
    rn2 = res.rezultat_net[2]

    if mc2 > pe2:
        motor_sig = f"formarea marjelor în perioada de raportare 2 este susținută preponderent de activitatea comercială (Marja comercială = {fmt_money(mc2)} lei, față de Producția exercițiului = {fmt_money(pe2)} lei)"
    elif pe2 > 0 and mc2 > 0:
        motor_sig = f"baza generatoare de valoare în perioada de raportare 2 îmbină Producția exercițiului ({fmt_money(pe2)} lei) cu Marja comercială ({fmt_money(mc2)} lei)"
    else:
        motor_sig = f"activitatea de bază din perioada de raportare 2 se sprijină pe Producția exercițiului ({fmt_money(pe2)} lei)"

    s13_end = pick_variant(
        bal,
        "s13_end",
        [
            f"La treptele finale ale cascadei din perioada de raportare 2, Rezultatul exploatării este de {fmt_money(re2)} lei, Rezultatul curent atinge {fmt_money(rc2)} lei, iar Rezultatul net al exercițiului se cifrează la {fmt_money(rn2)} lei.",
            f"După deducerea amortizărilor și integrarea fluxurilor financiare și fiscale în perioada de raportare 2, cascada se închide cu un Rezultat al exploatării de {fmt_money(re2)} lei, un Rezultat curent de {fmt_money(rc2)} lei și un Profit net de {fmt_money(rn2)} lei.",
            f"Translatarea EBE ({fmt_money(ebe2)} lei) către palierele contabile finale din perioada de raportare 2 conduce la un Rezultat din exploatare de {fmt_money(re2)} lei, un Rezultat curent de {fmt_money(rc2)} lei și un Rezultat net de {fmt_money(rn2)} lei.",
        ],
    )

    add_paragraph(
        doc,
        f"Interpretarea cascadei SIG: Analiza treptelor de acumulare arată că {motor_sig}; după scăderea consumurilor intermediare "
        f"provenite de la terți ({fmt_money(res.consumuri_terti[2])} lei în perioada de raportare 2), {traj_va}, în timp ce {traj_ebe}. {s13_end}",
    )
