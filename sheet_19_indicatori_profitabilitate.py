#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Capitolul 19: Analiza Indicatorilor de Profitabilitate și Rentabilitate.
Evaluează performanța relativă a societății prin șase rate de rentabilitate:
Marja brută din vânzări, Marja netă din vânzări, Rentabilitatea resurselor consumate,
Rentabilitatea economică (ROA), Rentabilitatea financiară (ROE) și Rentabilitatea comercială (ROS).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import List, Tuple

from docx import Document

from core_balanta import (
    BalantaData,
    add_heading,
    add_paragraph,
    calc_delta,
    calc_index,
    fmt_num,
    fmt_pct,
    fmt_ratio_pct,
    make_table,
    pick_variant,
    safe_div,
)
from sheet_02_cpp import CPPResult
from sheet_03_bilant_financiar import BilantFinanciarResult


@dataclass
class ProfitBlock:
    title: str
    num_name: str
    den_name: str
    rate_name: str
    formula_str: str
    num_vals: Tuple[float, float, float]
    den_vals: Tuple[float, float, float]
    rate_pct_vals: Tuple[float, float, float]


@dataclass
class IndicatoriProfitabilitateResult:
    marja_bruta: Tuple[float, float, float]
    marja_neta: Tuple[float, float, float]
    rent_resurse: Tuple[float, float, float]
    roa: Tuple[float, float, float]
    roe: Tuple[float, float, float]
    ros: Tuple[float, float, float]
    blocks: List[ProfitBlock] = field(default_factory=list)


def compute_sheet_19_indicatori_profitabilitate(
    bf: BilantFinanciarResult, cpp: CPPResult
) -> IndicatoriProfitabilitateResult:
    pb = cpp.rezultat_brut
    pn = cpp.rezultat_net
    ca = cpp.cifra_afaceri
    ct = cpp.cheltuieli_totale
    ta = bf.total_activ
    cp = bf.capitaluri_proprii

    mb = tuple(safe_div(pb[p], ca[p]) * 100.0 for p in range(3))
    mn = tuple(safe_div(pn[p], ca[p]) * 100.0 for p in range(3))
    rrc = tuple(safe_div(pn[p], ct[p]) * 100.0 for p in range(3))
    roa = tuple(safe_div(pn[p], ta[p]) * 100.0 for p in range(3))
    roe = tuple(safe_div(pn[p], cp[p]) * 100.0 for p in range(3))
    ros = tuple(safe_div(pn[p], ca[p]) * 100.0 for p in range(3))

    blocks = [
        ProfitBlock("Marja brută din vânzări", "Profit brut", "Cifra de afaceri", "1. Marja brută din vânzări", "(Profit brut / CA) × 100", pb, ca, mb),
        ProfitBlock("Marja netă din vânzări", "Profit net", "Cifra de afaceri", "2. Marja netă din vânzări", "(Profit net / CA) × 100", pn, ca, mn),
        ProfitBlock("Rentabilitatea resurselor consumate", "Profit net", "Cheltuieli totale", "3. Rentabilitatea resurselor consumate", "(Profit net / Cheltuieli totale) × 100", pn, ct, rrc),
        ProfitBlock("Rentabilitatea economică (ROA)", "Profit net", "Activ total", "4. Rentabilitatea economică (ROA)", "(Profit net / Activ total) × 100", pn, ta, roa),
        ProfitBlock("Rentabilitatea financiară (ROE)", "Profit net", "Capital propriu", "5. Rentabilitatea financiară (ROE)", "(Profit net / Capital propriu) × 100", pn, cp, roe),
        ProfitBlock("Rentabilitatea comercială (ROS)", "Profit net", "Cifra de afaceri", "6. Rentabilitatea comercială (ROS)", "(Profit net / CA) × 100", pn, ca, ros),
    ]

    return IndicatoriProfitabilitateResult(
        marja_bruta=mb,
        marja_neta=mn,
        rent_resurse=rrc,
        roa=roa,
        roe=roe,
        ros=ros,
        blocks=blocks,
    )


def render_sheet_19_indicatori_profitabilitate(
    doc: Document, bal: BalantaData, res: IndicatoriProfitabilitateResult
) -> None:
    """Generează secțiunea Indicatorilor de Profitabilitate într-un tabel unificat."""
    add_heading(doc, "19. Analiza Indicatorilor de Profitabilitate și Rentabilitate", level=1)

    headers = [
        "Indicator de Profitabilitate și Rentabilitate",
        "Formulă de Calcul",
        "Perioada de raportare 0",
        "Perioada de raportare 1",
        "Perioada de raportare 2",
        "Δ (P1 - P0)",
        "Δ (P2 - P1)",
        "I (P2 / P1)",
    ]
    rows = []
    for b in res.blocks:
        d = calc_delta(b.rate_pct_vals)
        i = calc_index(b.rate_pct_vals)
        rows.append([
            b.rate_name,
            b.formula_str,
            fmt_pct(b.rate_pct_vals[0]),
            fmt_pct(b.rate_pct_vals[1]),
            fmt_pct(b.rate_pct_vals[2]),
            fmt_pct(d[0], plus_sign=True),
            fmt_pct(d[1], plus_sign=True),
            fmt_ratio_pct(i[1]),
        ])

    make_table(
        doc,
        headers=headers,
        rows=rows,
        col_widths=[4.8, 2.8, 1.8, 1.8, 1.8, 1.6, 1.6, 1.6],
        total_rows=[3, 4, 5],
        font_size=7.8,
    )

    d_roa = res.roa[2] - res.roa[0]
    d_roe = res.roe[2] - res.roe[0]
    s19_msg = pick_variant(
        bal,
        "s19_prof",
        [
            f"Sinteză privind rentabilitatea: Rentabilitatea economică a activelor (ROA) evoluează de la {fmt_pct(res.roa[0])} în perioada de raportare 0 la {fmt_pct(res.roa[1])} în perioada de raportare 1 și {fmt_pct(res.roa[2])} în perioada de raportare 2 ({fmt_pct(d_roa, plus_sign=True)} pp față de perioada 0), în timp ce rentabilitatea financiară (ROE) înregistrează {fmt_pct(res.roe[0])}, {fmt_pct(res.roe[1])} și {fmt_pct(res.roe[2])} ({fmt_pct(d_roe, plus_sign=True)} pp). În perioada de raportare 2, la fiecare 100 lei de cheltuieli totale angajate societatea obține un rezultat net de {fmt_num(res.rent_resurse[2], 2)} lei (rentabilitatea resurselor consumate = {fmt_pct(res.rent_resurse[2])}, rentabilitatea comercială ROS = {fmt_pct(res.ros[2])}).",
            f"Evaluarea ratelor de eficiență: În perioada de raportare 2, randamentul activelor totale (ROA) este de {fmt_pct(res.roa[2])} (față de {fmt_pct(res.roa[0])} în perioada 0), remunerarea capitalurilor proprii (ROE) atinge {fmt_pct(res.roe[2])} (față de {fmt_pct(res.roe[0])} în perioada 0), iar marja comercială operațională (ROS) se situează la {fmt_pct(res.ros[2])}, cu o rentabilitate a resurselor consumate de {fmt_pct(res.rent_resurse[2])} și o marjă netă din vânzări de {fmt_pct(res.marja_neta[2])}.",
            f"Diagnostic de profitabilitate: Tabloul ratelor de rentabilitate din perioada de raportare 2 evidențiază o rentabilitate economică ROA = {fmt_pct(res.roa[2])}, o rentabilitate financiară ROE = {fmt_pct(res.roe[2])}, o rentabilitate comercială ROS = {fmt_pct(res.ros[2])} și o marjă brută din vânzări de {fmt_pct(res.marja_bruta[2])}, reflectând eficiența transformării resurselor consumate ({fmt_pct(res.rent_resurse[2])}) în profit net.",
        ],
    )
    add_paragraph(doc, s19_msg)
